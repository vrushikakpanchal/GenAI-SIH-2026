"""Document parsing for the Social Media Content Engine.

Text is extracted directly whenever possible. If a PDF has no text layer,
the parser automatically falls back to local Tesseract OCR.
"""

from __future__ import annotations

import io
import json
import os
from pathlib import Path
from typing import Optional, Union

import fitz
from docx import Document

try:
    from PIL import Image
    import pytesseract
except ImportError:
    Image = None
    pytesseract = None


ALLOWED_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".txt",
    ".json",
}

SourcePath = Union[str, os.PathLike[str]]

OCR_DPI = int(
    os.getenv(
        "SMC_OCR_DPI",
        "180",
    )
)

OCR_LANGUAGE = os.getenv(
    "SMC_OCR_LANGUAGE",
    "eng",
)

OCR_PSM = os.getenv(
    "SMC_OCR_PSM",
    "6",
)


if pytesseract is not None and os.getenv("TESSERACT_CMD"):
    pytesseract.pytesseract.tesseract_cmd = os.environ[
        "TESSERACT_CMD"
    ]


class DocumentParseError(ValueError):
    """Raised when a source document cannot be read."""


def clean_text(text: str) -> str:
    """Remove empty lines and trailing whitespace."""

    if not text:
        return ""

    paragraphs: list[str] = []

    for raw_line in (
        text.replace("\r\n", "\n")
        .replace("\r", "\n")
        .split("\n")
    ):
        line = " ".join(raw_line.split())

        if line:
            paragraphs.append(line)

    return "\n".join(paragraphs).strip()


def _ocr_pdf_pages(
    document: fitz.Document,
) -> str:
    """OCR every page of a scanned PDF with local Tesseract."""

    if pytesseract is None or Image is None:
        raise DocumentParseError(
            "This PDF appears to be scanned. Install the Python OCR "
            "packages and the Tesseract application, then try again."
        )

    scale = OCR_DPI / 72
    matrix = fitz.Matrix(scale, scale)
    pages: list[str] = []

    try:
        for page_number, page in enumerate(
            document,
            start=1,
        ):
            pixmap = page.get_pixmap(
                matrix=matrix,
                alpha=False,
            )

            image = Image.open(
                io.BytesIO(
                    pixmap.tobytes("png")
                )
            ).convert("RGB")

            text = pytesseract.image_to_string(
                image,
                lang=OCR_LANGUAGE,
                config=f"--psm {OCR_PSM}",
            )

            if text.strip():
                pages.append(
                    f"[Page {page_number}]\n{text}"
                )

    except RuntimeError as exc:
        if exc.__class__.__name__ == "TesseractNotFoundError":
            raise DocumentParseError(
                "This PDF appears to be scanned, but Tesseract "
                "was not found. Install Tesseract OCR or set "
                "TESSERACT_CMD to tesseract.exe."
            ) from exc

        raise DocumentParseError(
            f"Tesseract OCR failed: {exc}"
        ) from exc

    except OSError as exc:
        raise DocumentParseError(
            f"Could not render a PDF page for OCR: {exc}"
        ) from exc

    return clean_text(
        "\n".join(pages)
    )


def extract_from_pdf(
    file_path: SourcePath,
) -> str:
    """
    Extract a PDF text layer.

    If the PDF has no selectable text, OCR is used automatically.
    """

    try:
        with fitz.open(str(file_path)) as document:
            pages = [
                page.get_text("text")
                for page in document
            ]

            text = clean_text(
                "\n".join(pages)
            )

            if text:
                return text

            return _ocr_pdf_pages(document)

    except (
        fitz.FileDataError,
        OSError,
    ) as exc:
        raise DocumentParseError(
            "The PDF could not be opened. Check that it is "
            "a valid, readable PDF."
        ) from exc


def extract_from_docx(
    file_path: SourcePath,
) -> str:
    """Extract paragraphs and simple table text from DOCX."""

    try:
        document = Document(str(file_path))

    except (
        OSError,
        ValueError,
    ) as exc:
        raise DocumentParseError(
            "The DOCX could not be opened. Check that it is "
            "a valid Word document."
        ) from exc

    parts = [
        paragraph.text
        for paragraph in document.paragraphs
    ]

    for table in document.tables:
        for row in table.rows:
            parts.append(
                " | ".join(
                    cell.text
                    for cell in row.cells
                )
            )

    return clean_text(
        "\n".join(parts)
    )


def extract_from_txt(
    file_path: SourcePath,
) -> str:
    """Read a UTF-8 text file."""

    try:
        with open(
            file_path,
            "r",
            encoding="utf-8-sig",
        ) as source_file:
            return clean_text(
                source_file.read()
            )

    except UnicodeDecodeError as exc:
        raise DocumentParseError(
            "The TXT file is not UTF-8 encoded. "
            "Save it as UTF-8 and try again."
        ) from exc

    except OSError as exc:
        raise DocumentParseError(
            "The TXT file could not be read."
        ) from exc


def extract_from_json(
    file_path: SourcePath,
) -> str:
    """Convert JSON into readable text."""

    try:
        with open(
            file_path,
            "r",
            encoding="utf-8-sig",
        ) as source_file:
            data = json.load(source_file)

    except json.JSONDecodeError as exc:
        raise DocumentParseError(
            "The JSON file is not valid JSON."
        ) from exc

    except (
        OSError,
        UnicodeDecodeError,
    ) as exc:
        raise DocumentParseError(
            "The JSON file could not be read."
        ) from exc

    return clean_text(
        json.dumps(
            data,
            indent=2,
            ensure_ascii=False,
        )
    )


def extract_text(
    file_path: SourcePath,
) -> str:
    """Detect a supported file type and extract its text."""

    extension = Path(file_path).suffix.lower()

    if extension not in ALLOWED_EXTENSIONS:
        supported = ", ".join(
            sorted(ALLOWED_EXTENSIONS)
        )

        raise DocumentParseError(
            f"Unsupported file type. Please upload one of: "
            f"{supported}."
        )

    extractors = {
        ".pdf": extract_from_pdf,
        ".docx": extract_from_docx,
        ".txt": extract_from_txt,
        ".json": extract_from_json,
    }

    text = extractors[extension](file_path)

    if not text:
        if extension == ".pdf":
            raise DocumentParseError(
                "No text was found in this PDF, including after OCR. "
                "Check the scan quality or paste the text instead."
            )

        raise DocumentParseError(
            "The document contains no readable text."
        )

    return text


def parse_document(
    file_path: Optional[SourcePath] = None,
    pasted_text: Optional[str] = None,
) -> str:
    """Parse pasted text first, otherwise parse the uploaded file."""

    if pasted_text and pasted_text.strip():
        text = clean_text(pasted_text)

        if text:
            return text

    if file_path:
        return extract_text(file_path)

    raise DocumentParseError(
        "Please upload a file or paste source text."
    )


parse_source = parse_document